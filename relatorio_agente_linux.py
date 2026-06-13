#!/usr/bin/env python3
"""
Agente Automatico - Suporte Cismetro (Linux/VPS)
=================================================
1. Garante que o whatsapp-bridge esta rodando
2. Busca mensagens do grupo Suporte Cismetro (periodo configuravel)
3. Gera relatorio Excel com mensagens + analise de incidentes
4. Salva em /opt/cismetro/relatorios/ (acessivel via SFTP)

Uso:
    python3 relatorio_agente.py          # ultimos 30 dias
    python3 relatorio_agente.py 60       # ultimos 60 dias
    python3 relatorio_agente.py 7        # ultima semana
"""

import os
import sys
import time
import subprocess
import sqlite3
import signal
import logging
from datetime import datetime, timedelta
from collections import Counter
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    subprocess.run([sys.executable, "-m", "pip", "install", "openpyxl", "-q"])
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

try:
    import requests
except ImportError:
    subprocess.run([sys.executable, "-m", "pip", "install", "requests", "-q"])
    import requests

# ── Configuracoes ─────────────────────────────────────────────────────────────
BASE_DIR    = Path("/opt/cismetro")
BRIDGE_BIN  = BASE_DIR / "whatsapp-bridge"
BRIDGE_DIR  = BASE_DIR / "bridge"
DB_PATH     = BASE_DIR / "store" / "messages.db"
OUTPUT_DIR  = BASE_DIR / "relatorios"           # pasta SFTP
LOG_FILE    = BASE_DIR / "agente.log"
BRIDGE_URL  = "http://localhost:8080/api"

CHAT_JID    = "120363418632425656@g.us"
CHAT_NAME   = "Suporte Cismetro"
DIAS_PADRAO = 30

# ── Logging ───────────────────────────────────────────────────────────────────
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[
        logging.FileHandler(LOG_FILE, encoding="utf-8"),
        logging.StreamHandler(sys.stdout),
    ]
)
log = logging.getLogger("cismetro")

# ── Helpers de estilo Excel ───────────────────────────────────────────────────
def make_fill(hex_color): return PatternFill("solid", fgColor=hex_color)
def make_font(bold=False, color="000000", size=9):
    return Font(bold=bold, color=color, size=size, name="Calibri")

THIN   = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT   = Alignment(horizontal="left",   vertical="top",    wrap_text=True)

# ── 1. Bridge ─────────────────────────────────────────────────────────────────
def bridge_alive():
    try:
        r = requests.get(f"{BRIDGE_URL}/chats", timeout=3)
        return r.status_code < 500
    except Exception:
        return False

def start_bridge():
    if bridge_alive():
        log.info("Bridge ja esta rodando.")
        return None

    if not BRIDGE_BIN.exists():
        log.error(f"Binario nao encontrado: {BRIDGE_BIN}")
        log.error("Execute: cd /opt/cismetro && go build -o whatsapp-bridge .")
        sys.exit(1)

    log.info("Iniciando whatsapp-bridge...")
    proc = subprocess.Popen(
        [str(BRIDGE_BIN)],
        cwd=str(BRIDGE_DIR),
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
        preexec_fn=os.setsid
    )

    for i in range(30):
        time.sleep(1)
        if bridge_alive():
            log.info(f"Bridge pronto em {i+1}s (PID {proc.pid}).")
            return proc
        log.debug(f"Aguardando bridge... {i+1}s")

    log.warning("Bridge nao respondeu em 30s. Verifique a sessao do WhatsApp.")
    return proc

# ── 2. Busca mensagens ────────────────────────────────────────────────────────
def buscar_mensagens(dias: int):
    if not DB_PATH.exists():
        log.error(f"Banco nao encontrado: {DB_PATH}")
        log.error("Certifique-se de que /opt/cismetro/store/messages.db existe.")
        sys.exit(1)

    after = datetime.now() - timedelta(days=dias)
    conn  = sqlite3.connect(str(DB_PATH))
    conn.row_factory = sqlite3.Row
    cur   = conn.cursor()

    cur.execute("""
        SELECT
            m.id,
            m.timestamp,
            m.sender,
            m.content,
            m.is_from_me,
            COALESCE(c.name, m.sender) AS sender_name
        FROM messages m
        LEFT JOIN chats c ON c.jid = m.sender
        WHERE m.chat_jid = ?
          AND m.timestamp >= ?
          AND m.content IS NOT NULL
          AND m.content != ''
        ORDER BY m.timestamp ASC
    """, (CHAT_JID, after.isoformat()))

    rows = cur.fetchall()
    conn.close()
    log.info(f"{len(rows)} mensagens encontradas (ultimos {dias} dias).")
    return rows

# ── 3. Classificacao ──────────────────────────────────────────────────────────
KEYWORDS_INDISPONIBILIDADE = [
    "sem sistema", "sistema parado", "sistema offline", "sistema fora",
    "inoperante", "caiu o sistema", "sistema caiu", "sem acesso",
    "fora do ar", "sistema muito lento", "sistema oscilando", "nao carrega",
    "nao carrega", "off "
]
KEYWORDS_BUG = [
    "nao encontra", "nao aparece", "nao funciona", "nao esta", "erro",
    "bug", "falha", "problema", "parou", "caiu", "offline",
    "duplicado", "duplicata", "fantasma", "divergencia", "divergente",
    "nao consegue", "travando", "trava", "lento", "nao abre", "nao gera",
    "zerado", "bloqueado", "nao mostra", "sumiu", "desapareceu",
    "inconsistencia", "valor errado", "vdrl", "nao localiza", "nao carrega"
]
KEYWORDS_SUPORTE = [
    "vamos analisar", "iremos analisar", "so um minuto",
    "estamos verificando", "pode testar", "ja liberamos", "pronto",
    "deu certo", "resolvido", "foi corrigido", "liberado", "liberamos"
]

def normalize(text: str) -> str:
    import unicodedata
    return unicodedata.normalize("NFD", text).encode("ascii", "ignore").decode().lower()

def classify(text: str) -> str:
    t = normalize(text)
    if any(k in t for k in KEYWORDS_INDISPONIBILIDADE):
        return "Indisponibilidade"
    if any(k in t for k in KEYWORDS_BUG):
        return "Bug / Problema"
    if any(k in t for k in KEYWORDS_SUPORTE):
        return "Resposta Suporte"
    return "Geral"

CATEGORIA_CORES = {
    "Indisponibilidade": ("FFC7CE", "9C0006"),
    "Bug / Problema":    ("FFE699", "833C00"),
    "Resposta Suporte":  ("C6EFCE", "375623"),
    "Geral":             ("F2F2F2", "444444"),
}

# ── 4. Gera Excel ─────────────────────────────────────────────────────────────
def gerar_excel(rows, dias: int) -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    ts     = datetime.now().strftime("%Y%m%d_%H%M")
    output = OUTPUT_DIR / f"Cismetro_{ts}.xlsx"

    wb = Workbook()

    # ── Aba 1: Mensagens ──────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Mensagens"
    ws.freeze_panes = "A2"

    headers = ["#", "Data/Hora", "Remetente", "Mensagem", "Categoria", "De Mim?"]
    widths  = [5, 18, 22, 80, 22, 8]
    for col, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = make_fill("1F3964")
        cell.font = make_font(bold=True, color="FFFFFF", size=10)
        cell.alignment = CENTER
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[1].height = 22

    data_rows = []
    for msg in rows:
        ts_val = msg["timestamp"]
        try:
            dt = datetime.fromisoformat(ts_val) if isinstance(ts_val, str) \
                 else datetime.fromtimestamp(float(ts_val))
        except Exception:
            dt = datetime.now()
        content   = msg["content"] or ""
        sender    = msg["sender_name"] or msg["sender"] or "?"
        from_me   = "Sim" if msg["is_from_me"] else "Nao"
        categoria = classify(content)
        data_rows.append((dt, sender, content, categoria, from_me))

    for i, (dt, sender, content, categoria, from_me) in enumerate(data_rows, 1):
        bg, fg   = CATEGORIA_CORES.get(categoria, ("FFFFFF", "000000"))
        row_fill = make_fill(bg)
        vals = [i, dt.strftime("%d/%m/%Y %H:%M"), sender, content, categoria, from_me]
        aligns = [CENTER, CENTER, LEFT, LEFT, CENTER, CENTER]
        for col, (val, aln) in enumerate(zip(vals, aligns), 1):
            cell = ws.cell(row=i+1, column=col, value=val)
            cell.fill = row_fill
            cell.font = make_font(bold=(col == 5), color=fg if col == 5 else "000000")
            cell.alignment = aln
            cell.border = BORDER
        ws.row_dimensions[i+1].height = 55

    ws.auto_filter.ref = f"A1:F{len(data_rows)+1}"

    # ── Aba 2: Resumo ─────────────────────────────────────────────────────────
    ws2 = wb.create_sheet("Resumo")
    periodo_ini = (datetime.now() - timedelta(days=dias)).strftime("%d/%m/%Y")
    periodo_fim = datetime.now().strftime("%d/%m/%Y")

    ws2["B2"] = f"RELATORIO AUTOMATICO - {CHAT_NAME.upper()}"
    ws2["B2"].font = Font(bold=True, size=14, color="1F3964", name="Calibri")
    ws2.merge_cells("B2:F2")

    ws2["B3"] = (f"Periodo: {periodo_ini} a {periodo_fim}  |  "
                 f"Total: {len(data_rows)} mensagens  |  "
                 f"Gerado: {datetime.now().strftime('%d/%m/%Y %H:%M')}")
    ws2["B3"].font = Font(italic=True, size=9, color="595959", name="Calibri")
    ws2.merge_cells("B3:F3")

    cat_count = Counter(r[3] for r in data_rows)
    ws2["B5"] = "MENSAGENS POR CATEGORIA"
    ws2["B5"].font = make_font(bold=True, color="FFFFFF", size=11)
    ws2["B5"].fill = make_fill("1F3964")
    ws2.merge_cells("B5:D5")

    for i, cat in enumerate(["Indisponibilidade","Bug / Problema","Resposta Suporte","Geral"], 6):
        bg, fg = CATEGORIA_CORES[cat]
        cnt    = cat_count.get(cat, 0)
        pct    = f"{cnt/len(data_rows)*100:.1f}%" if data_rows else "0%"
        for col, val in [(2, cat), (3, cnt), (4, pct)]:
            c = ws2.cell(row=i, column=col, value=val)
            c.fill = make_fill(bg)
            c.font = make_font(bold=True, color=fg)
            c.border = BORDER
            c.alignment = LEFT if col == 2 else CENTER

    sender_count = Counter(r[1] for r in data_rows if r[4] == "Nao")
    ws2["B12"] = "TOP 10 REMETENTES (externos)"
    ws2["B12"].font = make_font(bold=True, color="FFFFFF", size=11)
    ws2["B12"].fill = make_fill("1F3964")
    ws2.merge_cells("B12:D12")

    for i, (name, cnt) in enumerate(sender_count.most_common(10), 13):
        ws2.cell(row=i, column=2, value=name).font = make_font(size=9)
        ws2.cell(row=i, column=3, value=cnt).font  = make_font(bold=True, size=9)
        ws2.cell(row=i, column=3).alignment = CENTER
        ws2.cell(row=i, column=4, value="=" * min(cnt, 30)).font = \
            Font(color="1F3964", size=9, name="Courier New")
        ws2.cell(row=i, column=2).border = ws2.cell(row=i, column=3).border = BORDER

    inc_por_dia = Counter(
        r[0].strftime("%d/%m/%Y") for r in data_rows
        if r[3] in ("Indisponibilidade", "Bug / Problema")
    )
    ws2["F5"] = "DIAS COM MAIS INCIDENTES"
    ws2["F5"].font = make_font(bold=True, color="FFFFFF", size=11)
    ws2["F5"].fill = make_fill("1F3964")
    ws2.merge_cells("F5:H5")

    for i, (day, cnt) in enumerate(inc_por_dia.most_common(10), 6):
        ws2.cell(row=i, column=6, value=day).font = make_font(size=9)
        ws2.cell(row=i, column=7, value=cnt).font = make_font(bold=True, size=9)
        ws2.cell(row=i, column=7).alignment = CENTER
        ws2.cell(row=i, column=6).border = ws2.cell(row=i, column=7).border = BORDER

    for col, w in [("B",28),("C",10),("D",25),("E",4),("F",14),("G",8)]:
        ws2.column_dimensions[col].width = w

    # ── Aba 3: Incidentes ─────────────────────────────────────────────────────
    ws3 = wb.create_sheet("Incidentes")
    ws3.freeze_panes = "A2"
    for col, (h, w) in enumerate(zip(
            ["#","Data/Hora","Remetente","Mensagem","Categoria"],
            [5, 18, 22, 90, 22]), 1):
        cell = ws3.cell(row=1, column=col, value=h)
        cell.fill = make_fill("C00000")
        cell.font = make_font(bold=True, color="FFFFFF", size=10)
        cell.alignment = CENTER
        cell.border = BORDER
        ws3.column_dimensions[get_column_letter(col)].width = w

    inc_data = [r for r in data_rows if r[3] in ("Indisponibilidade","Bug / Problema")]
    for i, (dt, sender, content, cat, _) in enumerate(inc_data, 1):
        bg, fg = CATEGORIA_CORES[cat]
        vals   = [i, dt.strftime("%d/%m/%Y %H:%M"), sender, content, cat]
        aligns = [CENTER, CENTER, LEFT, LEFT, CENTER]
        for col, (val, aln) in enumerate(zip(vals, aligns), 1):
            cell = ws3.cell(row=i+1, column=col, value=val)
            cell.fill = make_fill(bg)
            cell.font = make_font(bold=(col == 5), color=fg if col == 5 else "000000")
            cell.alignment = aln
            cell.border = BORDER
        ws3.row_dimensions[i+1].height = 55

    ws3.auto_filter.ref = f"A1:E{len(inc_data)+1}"

    wb.save(str(output))
    return output

# ── Main ───────────────────────────────────────────────────────────────────────
def main():
    dias = DIAS_PADRAO
    if len(sys.argv) > 1:
        try:
            dias = int(sys.argv[1])
        except ValueError:
            log.warning(f"Argumento invalido: {sys.argv[1]}. Usando {DIAS_PADRAO} dias.")

    log.info("=" * 50)
    log.info(f"AGENTE CISMETRO - Periodo: ultimos {dias} dias")
    log.info("=" * 50)

    start_bridge()
    rows  = buscar_mensagens(dias)
    log.info("Gerando relatorio Excel...")
    path  = gerar_excel(rows, dias)
    log.info(f"Relatorio salvo: {path}")

    # Ajusta permissoes para acesso SFTP
    os.chmod(str(path), 0o644)
    log.info("Concluido.")

if __name__ == "__main__":
    main()
